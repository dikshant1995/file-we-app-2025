import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\OneDrive\Desktop\company category"
master_path = os.path.join(folder, "ALL COMPANY IN ONE LIST.xlsx")

smfg_out_path = os.path.join(folder, "SMFG COM LIST.xlsx")
tata_out_path = os.path.join(folder, "TATA COM LIST.xlsx")

print("==========================================================================")
print("📌 EXTRACTING SMFG & TATA COMPANY LISTS FROM MASTER FILE")
print("==========================================================================\n")

if not os.path.exists(master_path):
    print(f"Error: Master file not found at {master_path}")
    sys.exit(1)

wb_master = openpyxl.load_workbook(master_path, data_only=True, read_only=True)
ws_master = wb_master['Sheet1']

header = None
bank_idx = None
cat_idx = None
comp_idx = None

smfg_rows = []
tata_rows = []

print("Reading master file records...", flush=True)
record_count = 0

for row in ws_master.iter_rows(values_only=True):
    if not any(row):
        continue
    row_str = [str(c).strip() if c is not None else "" for c in row]
    
    if header is None:
        header = row_str
        print(f"Header: {header}", flush=True)
        for i, h in enumerate(header):
            hl = h.lower()
            if 'bank' in hl:
                bank_idx = i
            elif 'cat' in hl:
                cat_idx = i
            elif 'company' in hl or 'name' in hl:
                comp_idx = i
        if bank_idx is None: bank_idx = 2
        if cat_idx is None: cat_idx = 1
        if comp_idx is None: comp_idx = 0
        continue
    
    record_count += 1
    bank_val = row_str[bank_idx] if bank_idx < len(row_str) else ""
    bank_val_upper = bank_val.upper()
    
    if "SMFG" in bank_val_upper:
        smfg_rows.append((row_str[comp_idx], row_str[cat_idx], bank_val))
    elif "TATA" in bank_val_upper:
        tata_rows.append((row_str[comp_idx], row_str[cat_idx], bank_val))

wb_master.close()

print(f"Read {record_count:,} master records.")
print(f"  • SMFG Records Extracted: {len(smfg_rows):,}")
print(f"  • Tata Records Extracted: {len(tata_rows):,}\n")

# Save SMFG COM LIST.xlsx
print(f"Writing {smfg_out_path}...", flush=True)
wb_smfg = openpyxl.Workbook()
ws_smfg = wb_smfg.active
ws_smfg.title = "SMFG Company List"
ws_smfg.append(["Sr.No", "Company Name", "Category", "Bank"])
for idx, (comp, cat, bname) in enumerate(smfg_rows, 1):
    ws_smfg.append([idx, comp, cat, bname])
wb_smfg.save(smfg_out_path)
wb_smfg.close()
print(f"✅ Saved SMFG COM LIST.xlsx successfully!")

# Save TATA COM LIST.xlsx
print(f"Writing {tata_out_path}...", flush=True)
wb_tata = openpyxl.Workbook()
ws_tata = wb_tata.active
ws_tata.title = "Tata Capital Company List"
ws_tata.append(["Sr.No", "Company Name", "Category", "Bank"])
for idx, (comp, cat, bname) in enumerate(tata_rows, 1):
    ws_tata.append([idx, comp, cat, bname])
wb_tata.save(tata_out_path)
wb_tata.close()
print(f"✅ Saved TATA COM LIST.xlsx successfully!")

print("\n==========================================================================")
print("🎉 EXTRACTION & SAVING COMPLETE!")
print("==========================================================================")

