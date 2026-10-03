import openpyxl
import os
import glob
import json

folder = r"C:\Users\user\OneDrive\Desktop\company category"
files = os.listdir(folder)

print(f"Total files found in folder: {len(files)}")
for f in files:
    print(f" - {f}")

report = {}

excel_files = [os.path.join(folder, f) for f in files if f.endswith('.xlsx')]

for fpath in excel_files:
    fname = os.path.basename(fpath)
    print(f"\n==================================================")
    print(f"Processing Excel: {fname}")
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True, read_only=True)
        sheets_summary = {}
        for sname in wb.sheetnames:
            ws = wb[sname]
            header = None
            row_count = 0
            sample_rows = []
            
            for r in ws.iter_rows(values_only=True):
                if not any(r):
                    continue
                row_count += 1
                clean_r = [str(cell).strip() if cell is not None else "" for cell in r]
                if header is None:
                    header = clean_r
                elif len(sample_rows) < 5:
                    sample_rows.append(clean_r)
            
            sheets_summary[sname] = {
                "rowCount": row_count,
                "header": header[:10] if header else [],
                "sample": sample_rows
            }
        wb.close()
        report[fname] = sheets_summary
        print(f"Finished {fname}. Sheets: {list(sheets_summary.keys())}")
    except Exception as e:
        print(f"Error processing {fname}: {e}")

# Save json report
with open(os.path.join(os.path.dirname(__file__), "excel_analysis.json"), "w", encoding="utf-8") as out:
    json.dump(report, out, indent=2)

print("\nSaved excel_analysis.json!")
